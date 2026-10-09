"""Prepare Qwen classification data without modifying the published source files.

Run: training/.qwen-venv/Scripts/python training/qwen/data.py
Outputs stay in gitignored artifacts/data. Spam means Suspicious, explicit
smishing means Scam, and ham means Low risk; this does not assert ham is safe.
Conflicting normalized labels and overlap with existing evaluation messages
are quarantined with source references. Holdouts are never used for fitting.
"""

import csv
import hashlib
import json
import logging
import re
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Literal

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
OUTPUT = Path(__file__).resolve().parent / "artifacts" / "data"
Verdict = Literal["scam", "suspicious", "low_risk"]


@dataclass(frozen=True)
class Message:
    text: str
    verdict: Verdict
    source: str
    row: int
    original_label: str


def fingerprint(text: str) -> str:
    normalized = re.sub(r"\s+", " ", re.sub(r"[0-9]", "0", text.lower())).strip()
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()


def family_fingerprint(text: str) -> str:
    """Group URL/number variants together for splitting without changing training text."""
    normalized = re.sub(r"(?:https?://|www\.)\S+", "URL", text.lower())
    normalized = re.sub(r"[0-9]+", "NUM", normalized)
    return fingerprint(normalized)


def source_verdict(label: str) -> Verdict:
    if label in {"scam", "smishing", "phishing", "fraud"}:
        return "scam"
    if label == "spam":
        return "suspicious"
    if label in {"legit", "ham", "ads", "notifs", "otp", "gov"}:
        return "low_risk"
    raise ValueError(f"Unsupported source label {label!r}; define its classification semantics explicitly")


def read_csv(path: Path) -> tuple[list[list[str]], list[int]]:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream, strict=True)
        header = next(reader)
        rows = [header]
        malformed: list[int] = []
        for index, row in enumerate(reader, 2):
            if len(row) != len(header):
                malformed.append(index)
            else:
                rows.append(row)
    return rows, malformed


def read_excel(path: Path) -> list[list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        if worksheet is None:
            raise ValueError(f"Workbook has no active sheet: {path}")
        rows = [["" if value is None else str(value) for value in row]
                for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    return rows


def load_messages(path: Path) -> tuple[list[Message], dict[str, int]]:
    rows, malformed = (read_excel(path), []) if path.suffix == ".xlsx" else read_csv(path)
    if not rows:
        raise ValueError(f"Dataset is empty: {path}")
    headers = [column.lower() for column in rows[0]]
    text_column = next((headers.index(name) for name in ("text", "message", "sms", "content")
                        if name in headers), None)
    label_column = next((headers.index(name) for name in ("label", "category", "type", "class")
                         if name in headers), None)
    if text_column is None or label_column is None:
        raise ValueError(f"Dataset requires text and label/category columns: {path}; headers={headers}")
    excluded: Counter[str] = Counter({"malformed": len(malformed)})
    messages: list[Message] = []
    source = path.relative_to(ROOT).as_posix()
    for index, row in enumerate(rows[1:], 2):
        # References for well-formed CSV records account for earlier excluded records.
        source_row = index
        for bad_row in malformed:
            if bad_row <= source_row:
                source_row += 1
        text = row[text_column].strip()
        if not text or text.upper() in {"<REDACTED>", "<BLANK>", "NAN"}:
            excluded["missing_message_content"] += 1
            continue
        if not any(character.isalpha() for character in text):
            excluded["no_linguistic_content"] += 1
            continue
        label = row[label_column].strip().lower()
        messages.append(Message(text, source_verdict(label), source, source_row, label))
    if sum(excluded.values()):
        logging.warning("Excluded unusable dataset records", extra={"source": source,
                        "exclusions": dict(excluded), "malformed_rows": malformed})
    return messages, dict(excluded)


def evaluation_fingerprints() -> set[str]:
    hashes: set[str] = set()
    for directory in ("dev", "test", "test_v2", "adversarial"):
        for path in sorted((ROOT / "data" / directory).glob("*.csv")):
            messages, _ = load_messages(path)
            hashes.update(family_fingerprint(message.text) for message in messages)
    path = ROOT / "data" / "llm_eval" / "cases.jsonl"
    for index, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        record = json.loads(line)
        text = record.get("message")
        if not isinstance(text, str) or not text:
            raise ValueError(f"LLM evaluation case has no message: {path}:{index}")
        hashes.add(family_fingerprint(text))
    return hashes


def split_messages(messages: list[Message]) -> dict[str, list[Message]]:
    splits: dict[str, list[Message]] = {"train": [], "tune": [], "validation": []}
    groups: dict[str, list[Message]] = defaultdict(list)
    for message in messages:
        groups[family_fingerprint(message.text)].append(message)
    for key in sorted(groups):
        group = sorted(groups[key], key=lambda message: fingerprint(message.text))
        bucket = int(hashlib.sha256(("42:" + key).encode()).hexdigest()[:8], 16) % 10
        split = "validation" if bucket == 0 else "tune" if bucket == 1 else "train"
        splits[split].extend(group)
    return splits


def main() -> None:
    paths = [*sorted((ROOT / "data" / "train").glob("*.csv")),
             ROOT / "data/raw/ph_spam/text-messages.csv",
             ROOT / "data/raw/tagalog_sms/tagalog-sms.xlsx",
             ROOT / "data/raw/mishra_soni/Dataset_5971.csv"]
    frozen = ROOT / "data/test_v2/test_v2.csv"
    frozen_hash = hashlib.sha256(frozen.read_bytes()).hexdigest()
    grouped: dict[str, list[Message]] = defaultdict(list)
    input_metadata: list[dict[str, str | int | dict[str, int]]] = []
    for path in paths:
        messages, exclusions = load_messages(path)
        input_metadata.append({"path": path.relative_to(ROOT).as_posix(),
                               "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                               "retained_rows": len(messages), "exclusions": exclusions})
        for message in messages:
            grouped[fingerprint(message.text)].append(message)
    protected = evaluation_fingerprints()
    conflicts: list[dict[str, str | list[dict[str, str | int]]]] = []
    overlap: list[str] = []
    unique: list[Message] = []
    for digest, group in sorted(grouped.items()):
        if len({message.verdict for message in group}) > 1:
            conflicts.append({"fingerprint": digest, "references": [
                {"source": message.source, "row": message.row, "label": message.original_label}
                for message in group]})
        elif family_fingerprint(group[0].text) in protected:
            overlap.append(digest)
        else:
            unique.append(group[0])
    splits = split_messages(unique)
    OUTPUT.mkdir(parents=True, exist_ok=True)
    for name, messages in splits.items():
        output = "".join(json.dumps(asdict(message), ensure_ascii=False) + "\n" for message in messages)
        (OUTPUT / f"{name}.jsonl").write_text(output, encoding="utf-8")
    report = {"seed": 42, "split_policy": "hash message families: 80% train, 10% tune, 10% validation",
              "label_policy": "spam=suspicious; smishing/scam=scam; ham/legit=low_risk",
              "inputs": input_metadata, "unique_before_quarantine": len(grouped),
              "duplicate_rows": sum(len(group) - 1 for group in grouped.values()),
              "conflict_groups": conflicts, "existing_evaluation_overlap": overlap,
              "splits": {name: dict(Counter(message.verdict for message in messages))
                         for name, messages in splits.items()},
              "split_hashes": {name: hashlib.sha256((OUTPUT / f"{name}.jsonl").read_bytes()).hexdigest()
                               for name in splits}, "frozen_test_sha256": frozen_hash,
              "limitations": ["AI-written synthetic labels still need team review",
                               "Same-inbox message-family splits are not a cross-inbox benchmark",
                               "Spam is not labeled as proven fraud",
                               "Normalized exact deduplication does not prove absence of near-duplicates"]}
    (OUTPUT / "manifest.json").write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    if hashlib.sha256(frozen.read_bytes()).hexdigest() != frozen_hash:
        raise RuntimeError("Frozen test file changed while preparing Qwen data")
    print(json.dumps({"output": str(OUTPUT), "conflict_groups": len(conflicts),
                      "protected_overlap": len(overlap), "splits": report["splits"]}, indent=2))


if __name__ == "__main__":
    logging.basicConfig(level=logging.WARNING, format="%(levelname)s %(message)s")
    main()
